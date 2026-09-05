import csv
from datetime import datetime, timezone

from openpyxl import load_workbook

from impurity_leads.contacts.models import ContactCandidate, ContactEvidence, RankedContact
from impurity_leads.export import export_csv, export_xlsx
from impurity_leads.models import Opportunity


def _opportunity():
    return Opportunity(
        company_name="Example Pharma",
        molecule_name="Drug X",
        score=80,
        priority="HOT",
        rationale="nitrosamine",
        likely_need_type=["nitrosamine/NDSRI"],
        decision_maker_role=["Analytical Development"],
        signal_fingerprints=["abc"],
    )


def _contacts():
    now = datetime(2026, 9, 4, tzinfo=timezone.utc)
    evidence = ContactEvidence(
        evidence_type="email",
        observed_value="jane.smith@example.com",
        source_type="company_site",
        source_url="https://example.com/team",
        evidence_text="Jane Smith",
        observed_at=now,
        fingerprint="fp1",
    )
    contact = ContactCandidate(
        company_name="Example Pharma",
        display_name="Jane Smith",
        normalized_name="jane smith",
        title="Director, Analytical Development",
        department="analytical",
        primary_email="jane.smith@example.com",
        email_status="VERIFIED_PUBLIC",
        role_fit_score=85,
        contact_confidence=1.0,
        last_verified_at=now,
        evidence=[evidence],
    )
    return [RankedContact(contact=contact, rank=1, rationale="target analytical leader")]


def test_exports_csv_and_xlsx(tmp_path):
    opportunity = _opportunity()
    mapping = {(opportunity.company_name, opportunity.molecule_name): _contacts()}
    csv_path = export_csv([opportunity], tmp_path / "opportunities.csv", mapping)
    xlsx_path = export_xlsx([opportunity], tmp_path / "opportunities.xlsx", mapping)
    assert csv_path.exists()
    assert xlsx_path.exists()

    with csv_path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    assert "Contact 1 Email Status" in rows[0]
    assert "jane.smith@example.com" in rows[1]
    assert "VERIFIED_PUBLIC" in rows[1]

    workbook = load_workbook(xlsx_path, read_only=True)
    row = next(workbook["Opportunities"].iter_rows(values_only=True))
    assert "Contact 1 Evidence URLs" in row
